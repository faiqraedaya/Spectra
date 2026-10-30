"""Excel and CSV reports with full traceability.

Every frequency in the summary can be traced through the parts count sheet
to the individual items, each with its drawing, page, position, source and
the line number or label its size came from.
"""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from .frequency import HOLE_BANDS, FrequencyTable, Results, calculate
from .geometry import pt_to_mm
from .project import Project


def write_xlsx(project: Project, table: FrequencyTable, path: str | Path, results: Results | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    res = results or calculate(project, table)
    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E5E7EB")

    def sheet(title: str, header: list[str], rows: list[list], formats: dict[int, str] | None = None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font = bold
            c.fill = head_fill
        for r in rows:
            ws.append(r)
        for col, fmt in (formats or {}).items():
            for cell in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                cell[0].number_format = fmt
        for i, h in enumerate(header, start=1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500] if i - 1 < len(r)]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
        ws.freeze_panes = "A2"
        return ws

    sci = "0.00E+00"
    summary = [[r.section] + [round(v, 12) for v in r.by_band] + [round(r.total, 12)] for r in res.sections]
    sheet("Frequencies", ["Section"] + [f"{b} (/yr)" for b in HOLE_BANDS] + ["Total (/yr)"], summary,
          {c: sci for c in range(2, 3 + len(HOLE_BANDS))})

    parts = []
    for r in res.sections:
        for p in r.parts:
            row = table.lookup(p.category, p.size_mm) if p.size_mm is not None else None
            freqs = [p.total * v for v in row.values] if row else [None] * len(HOLE_BANDS)
            parts.append([r.section, p.category, p.size_mm, p.explicit, p.implied, p.total, table.unit(p.category),
                          f"{row.min_mm:g}-{row.max_mm:g}" if row else "no size"] + freqs)
    sheet("Parts count", ["Section", "Category", "Size (mm)", "Counted", "Implied", "Total", "Basis", "Data row (mm)"]
          + [f"{b} (/yr)" for b in HOLE_BANDS], parts, {c: sci for c in range(9, 9 + len(HOLE_BANDS))})

    docs = {d.id: d for d in project.documents}
    tax = project.taxonomy
    items = []
    for i in sorted(project.items, key=lambda i: (docs[i.doc].name if i.doc in docs else "", i.page, i.box[1], i.box[0])):
        sec = project.section_of(i)
        x, y = (i.box[0] + i.box[2]) / 2, (i.box[1] + i.box[3]) / 2
        items.append([
            docs[i.doc].name if i.doc in docs else i.doc, i.page + 1, tax.name(i.cls), tax.get(i.cls).category or "not counted",
            project.effective_size(i), i.meta.get("size_from", "item" if i.size_mm is not None else "section/class default"),
            i.tag, i.count, sec.name if sec else "Unassigned", i.status, i.source, round(i.score, 3),
            round(pt_to_mm(x), 1), round(pt_to_mm(y), 1), i.meta.get("label", ""), i.id,
        ])
    sheet("Items", ["Drawing", "Page", "Class", "Category", "Size (mm)", "Size from", "Tag", "Count", "Section", "Status",
                    "Source", "Score", "X (mm)", "Y (mm)", "Markup label", "Item ID"], items)

    sheet("Issues", ["Item ID", "Issue"], [[i.item_id or "", i.message] for i in res.issues]
          + [["", f"Not counted (no frequency category): {k} x{v}"] for k, v in res.not_counted.items()])

    basis = [
        ["Project", project.name],
        ["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")],
        ["Frequency data", f"{table.meta.get('name', project.settings.frequency_table)}, {table.meta.get('edition', '')}"],
        ["Source", table.meta.get("source", "")],
        ["Included items", "accepted items only; proposals awaiting review are excluded"],
        ["Implied components", "; ".join(f"{c.name}: " + ", ".join(f"{n:g} x {k}" for k, n in c.implied.items())
                                         for c in tax if c.implied)],
        ["Class default sizes", "; ".join(f"{c.name}: {c.default_size_mm:g} mm" for c in tax if c.default_size_mm)],
        ["Drawings", "; ".join(f"{d.name} (sha256 {d.sha256[:12]})" for d in project.documents)],
    ]
    sheet("Basis", ["Item", "Value"], basis)
    del wb["Sheet"]
    wb.save(path)


def write_csv(results: Results, path: str | Path) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Section"] + [f"{b} (/yr)" for b in HOLE_BANDS] + ["Total (/yr)"])
        for r in results.sections:
            w.writerow([r.section] + [f"{v:.3e}" for v in r.by_band] + [f"{r.total:.3e}"])
